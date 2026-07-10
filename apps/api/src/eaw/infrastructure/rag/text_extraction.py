"""Full document text extraction for ingestion."""

from __future__ import annotations

import io
from pathlib import Path

from eaw.domain.common.errors import ValidationAppError


def extract_full_text(*, data: bytes, filename: str, content_type: str) -> str:
    ext = Path(filename).suffix.lower()

    if ext in {".txt", ".md", ".csv"} or content_type.startswith("text/"):
        return data.decode("utf-8", errors="replace")

    if ext == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            parts = [(p.extract_text() or "") for p in reader.pages]
            text = "\n".join(parts).strip()
            if not text:
                raise ValidationAppError(
                    "PDF has no extractable text (may be scanned — OCR in later phase)"
                )
            return text
        except ValidationAppError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise ValidationAppError(f"Failed to extract PDF text: {exc}") from exc

    if ext == ".docx":
        try:
            from docx import Document

            doc = Document(io.BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs if p.text)
        except Exception as exc:  # noqa: BLE001
            raise ValidationAppError(f"Failed to extract DOCX text: {exc}") from exc

    if ext == ".pptx":
        try:
            from pptx import Presentation

            prs = Presentation(io.BytesIO(data))
            parts: list[str] = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        parts.append(shape.text)
            return "\n".join(parts)
        except Exception as exc:  # noqa: BLE001
            raise ValidationAppError(f"Failed to extract PPTX text: {exc}") from exc

    raise ValidationAppError(f"Unsupported type for extraction: {ext}")
